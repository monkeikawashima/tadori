#!/usr/bin/env python3
"""調書（03_調書/*.json）を集計し、領域ごとの状態と未修正虚偽表示の合計を重要性と比べる。

使い方:
  python3 summarize.py --engagement <案件フォルダ> --today 2026-11-20 [--xlsx]

--xlsx を付けると 05_成果物 に未修正虚偽表示一覧を書き出す（版番号は自動で繰り上げ）。
終了コード: 0=全領域 pass かつ合計が手続上の重要性未満 / 1=それ以外（人の判断が要る）
"""
import argparse
import json
import re
import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import dump, load_engagement, to_decimal  # noqa: E402


def next_version(dir_: Path, stem: str) -> int:
    pat = re.compile(re.escape(stem) + r"_v(\d+)_")
    vs = [int(m.group(1)) for p in dir_.glob(f"{stem}_v*") if (m := pat.match(p.name))]
    return max(vs, default=0) + 1


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--engagement", required=True)
    ap.add_argument("--today", required=True)
    ap.add_argument("--xlsx", action="store_true")
    args = ap.parse_args()
    root = Path(args.engagement)
    eng = load_engagement(root)
    m = eng.get("materiality", {})
    overall = to_decimal(m.get("overall") or 0)
    performance = to_decimal(m.get("performance") or 0)
    trivial = to_decimal(m.get("trivial") or 0)

    areas, misstatements, escalations, blocked = [], [], [], []
    for p in sorted((root / "03_調書").glob("*.json")):
        if p.name.startswith("_") or p.name.endswith(".review.json"):
            continue
        wp = json.loads(p.read_text(encoding="utf-8"))
        aid = wp.get("area_id", p.stem)
        rounds = wp.get("rounds", [])
        areas.append({
            "area_id": aid, "area_name": wp.get("area_name", ""), "status": wp.get("status", "draft"),
            "confidence": wp.get("confidence", ""), "rounds": len(rounds),
            "last_score": rounds[-1].get("score") if rounds else None,
        })
        for ms in wp.get("misstatements", []):
            amt = to_decimal(ms.get("amount", 0))
            misstatements.append({**ms, "area_id": aid, "amount": amt, "below_trivial": bool(trivial and abs(amt) < trivial)})
        for e in wp.get("escalations", []):
            escalations.append({"area_id": aid, **e})
        for b in wp.get("blocked", []):
            blocked.append({"area_id": aid, **b})

    uncorrected = [x for x in misstatements if not x.get("corrected")]
    total = sum((abs(x["amount"]) for x in uncorrected if not x["below_trivial"]), Decimal(0))
    judgements = []
    if performance and total >= performance:
        judgements.append(f"未修正虚偽表示の合計 {total} が手続上の重要性 {performance} 以上")
    if overall and total >= overall:
        judgements.append(f"未修正虚偽表示の合計 {total} が全体の重要性 {overall} 以上")
    if overall and performance and performance <= total < overall:
        judgements.append("手続上の重要性と全体の重要性の間にある。追加手続の要否を人が判断する")
    not_pass = [a["area_id"] for a in areas if a["status"] != "pass"]

    result = {
        "areas": areas,
        "areas_not_pass": not_pass,
        "uncorrected_misstatements": uncorrected,
        "uncorrected_total_above_trivial": total,
        "materiality": {"overall": overall, "performance": performance, "trivial": trivial},
        "judgements_needed": judgements,
        "escalations": escalations,
        "blocked": blocked,
    }
    (root / "03_調書" / "_集計.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    if args.xlsx:
        import openpyxl  # noqa: PLC0415
        from openpyxl.styles import Font  # noqa: PLC0415
        stem = f'{eng.get("client", "")}_{eng.get("fiscal_period_end", "")[:7]}期_未修正虚偽表示一覧'
        out_dir = root / "05_成果物"
        v = next_version(out_dir, stem)
        path = out_dir / f'{stem}_v{v}_{args.today.replace("-", "")}.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "未修正虚偽表示"
        ws.append(["ID", "領域", "内容", "区分", "金額", "影響勘定", "証拠", "僅少"])
        for x in uncorrected:
            ws.append([x.get("id", ""), x["area_id"], x.get("description", ""), x.get("type", ""),
                       float(x["amount"]), "、".join(x.get("accounts", [])), "、".join(x.get("evidence", [])),
                       "僅少" if x["below_trivial"] else ""])
        n = ws.max_row
        ws.append([])
        ws.append(["", "", "合計（僅少を除く、絶対値）", "", f'=SUMPRODUCT(ABS(E2:E{n})*(H2:H{n}=""))'])
        ws.append(["", "", "手続上の重要性", "", float(performance)])
        ws.append(["", "", "全体の重要性", "", float(overall)])
        for c in ws[1]:
            c.font = Font(bold=True)
        for col, w in zip("ABCDEFGH", [12, 14, 48, 8, 16, 18, 24, 6]):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
            for c in row:
                c.number_format = "#,##0;[Red]-#,##0"
        ws.freeze_panes = "A2"
        wb.save(path)
        result["xlsx"] = str(path.relative_to(root))

    dump(result)
    return 0 if not not_pass and not judgements else 1


if __name__ == "__main__":
    sys.exit(main())
