#!/usr/bin/env python3
"""案件フォルダの作成・計画の承認と差戻し・案件台帳の更新。

サブコマンド:
  new      案件フォルダをテンプレから作り、案件台帳に1行足す
  approve  計画を承認済みにする（以後 plan.json と engagement.json は hooks で書込禁止）
  reopen   承認を取り消す。人の指示があったときだけ使う。理由は必須で、ログに残る
  status   案件台帳の状態列を更新する

例:
  python3 engagement.py new --root ~/監査案件 --client ABC商事 --period-end 2026-03-31 \
      --today 2026-10-27 [--prior ~/監査案件/ABC商事_2025-03期]
  python3 engagement.py approve --engagement <案件フォルダ> --by "川島" --today 2026-10-28
  python3 engagement.py status --engagement <案件フォルダ> --status "最終査閲待ち" --today 2026-11-20
"""
import argparse
import json
import re
import shutil
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import ENGAGEMENT_FILE, FOLDERS, dump, find_engagement_root, load_engagement  # noqa: E402

PLUGIN_ROOT = Path(__file__).resolve().parent.parent
LEDGER_HEADER = ["案件フォルダ", "クライアント", "決算期末", "種別", "状態", "計画承認", "最終更新"]


def safe(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "", s)


def ledger_path(root: Path) -> Path:
    return root / "案件台帳.xlsx"


def upsert_ledger(root: Path, row: dict) -> str:
    """案件台帳.xlsx（無ければ作る）の該当行を更新、無ければ追記する。openpyxl が無ければ CSV。"""
    try:
        import openpyxl  # noqa: PLC0415
        from openpyxl.styles import Font  # noqa: PLC0415
    except ImportError:
        import csv  # noqa: PLC0415
        p = root / "案件台帳.csv"
        rows = []
        if p.exists():
            with open(p, encoding="utf-8-sig", newline="") as f:
                rows = list(csv.DictReader(f))
        # xlsx と同じく、渡された値が空でない列だけを書き換える（状態の更新で計画承認などを消さない）
        for r in rows:
            if r["案件フォルダ"] == row["案件フォルダ"]:
                r.update({k: v for k, v in row.items() if v != ""})
                break
        else:
            rows.append(row)
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=LEDGER_HEADER)
            w.writeheader()
            w.writerows(rows)
        return p.name
    p = ledger_path(root)
    if p.exists():
        wb = openpyxl.load_workbook(p)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "案件台帳"
        ws.append(LEDGER_HEADER)
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, w in enumerate([28, 20, 12, 12, 16, 16, 12], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
    values = [row.get(h, "") for h in LEDGER_HEADER]
    for r in ws.iter_rows(min_row=2):
        if r[0].value == row["案件フォルダ"]:
            for c, v in zip(r, values):
                if v != "":
                    c.value = v
            break
    else:
        ws.append(values)
    ws.auto_filter.ref = ws.dimensions
    wb.save(p)
    return p.name


def cmd_new(args) -> None:
    root = Path(args.root).expanduser()
    root.mkdir(parents=True, exist_ok=True)
    period_label = args.period_end[:7]
    name = f"{safe(args.client)}_{period_label}期"
    eng_dir = root / name
    if eng_dir.exists():
        sys.exit(f"同名の案件フォルダが既にあります: {eng_dir}")
    for f in FOLDERS:
        (eng_dir / f).mkdir(parents=True)

    eng = json.loads((PLUGIN_ROOT / "templates" / "engagement.json").read_text(encoding="utf-8"))
    eng["client"] = args.client
    eng["fiscal_period_end"] = args.period_end
    if args.engagement_type:
        eng["engagement_type"] = args.engagement_type
    eng["folder_name"] = name
    (eng_dir / ENGAGEMENT_FILE).write_text(json.dumps(eng, ensure_ascii=False, indent=2), encoding="utf-8")

    carried = []
    if args.prior:
        prior = Path(args.prior).expanduser()
        for cand in [prior / "05_成果物" / "次期引継ぎ.md", prior / "00_案件情報" / "次期引継ぎ.md"]:
            if cand.exists():
                shutil.copy2(cand, eng_dir / "00_案件情報" / "前期引継ぎ.md")
                carried.append(str(cand))
                break
        prior_eng = prior / ENGAGEMENT_FILE
        if prior_eng.exists():
            pe = json.loads(prior_eng.read_text(encoding="utf-8"))
            eng["contacts"] = pe.get("contacts", eng["contacts"])
            eng["auditor_domains"] = pe.get("auditor_domains", [])
            eng["accounting_system"] = pe.get("accounting_system", "")
            (eng_dir / ENGAGEMENT_FILE).write_text(json.dumps(eng, ensure_ascii=False, indent=2), encoding="utf-8")
            carried.append("contacts / auditor_domains / accounting_system を前期から引継ぎ（計画承認で要確認）")

    readme = (PLUGIN_ROOT / "templates" / "案件フォルダの使い方.md")
    if readme.exists():
        shutil.copy2(readme, eng_dir / "案件フォルダの使い方.md")

    ledger = upsert_ledger(root, {
        "案件フォルダ": name, "クライアント": args.client, "決算期末": args.period_end,
        "種別": args.engagement_type or "", "状態": "計画中", "計画承認": "", "最終更新": args.today,
    })
    dump({"engagement": str(eng_dir), "ledger": ledger, "carried_over": carried})


def engagement_dir(arg: str) -> Path:
    """--engagement を絶対パスの案件フォルダにする。

    「.」のような相対パスのままだと親フォルダと案件フォルダ名が取れず、案件台帳を
    案件フォルダの中に作ってしまう。案件フォルダの中のサブフォルダを渡されても、親をたどって見つける。
    """
    root = find_engagement_root(Path(arg).expanduser())
    if root is None:
        sys.exit(f"案件フォルダが見つかりません（00_案件情報/engagement.json がない）: {arg}")
    return root


def cmd_approve(args) -> None:
    root = engagement_dir(args.engagement)
    eng = load_engagement(root)
    plan = root / "00_案件情報" / "plan.json"
    if not plan.exists():
        sys.exit("00_案件情報/plan.json がありません。先に計画を作ってください")
    m = eng.get("materiality", {})
    if not m.get("overall") or not m.get("performance"):
        sys.exit("重要性（overall と performance）が未設定です。人が決めた値を engagement.json に入れてから承認してください")
    if not eng.get("contacts") or not all(c.get("email") for c in eng["contacts"]):
        sys.exit("依頼先（contacts）のメールアドレスが空です")
    try:
        date.fromisoformat(eng.get("evidence_deadline") or "")
    except ValueError:
        sys.exit("資料の提出期限（evidence_deadline、YYYY-MM-DD）が未設定です。依頼メールの期限と督促の基準に使います")
    if not (eng.get("mail") or {}).get("signature", "").strip():
        sys.exit("依頼メールの署名（mail.signature）が未設定です")
    eng["plan_status"] = "approved"
    eng["plan_approved_by"] = args.by
    eng["plan_approved_at"] = args.today
    (root / ENGAGEMENT_FILE).write_text(json.dumps(eng, ensure_ascii=False, indent=2), encoding="utf-8")
    log = root / "06_ログ" / "承認記録.jsonl"
    with open(log, "a", encoding="utf-8") as f:
        f.write(json.dumps({"ts": args.today, "action": "approve", "by": args.by}, ensure_ascii=False) + "\n")
    upsert_ledger(root.parent, {"案件フォルダ": root.name, "状態": "実行中", "計画承認": f"{args.today} {args.by}", "最終更新": args.today})
    dump({"plan_status": "approved", "by": args.by, "at": args.today})


def cmd_reopen(args) -> None:
    if not args.reason:
        sys.exit("--reason が必要です")
    root = engagement_dir(args.engagement)
    eng = load_engagement(root)
    eng["plan_status"] = "draft"
    (root / ENGAGEMENT_FILE).write_text(json.dumps(eng, ensure_ascii=False, indent=2), encoding="utf-8")
    with open(root / "06_ログ" / "承認記録.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps({"ts": args.today, "action": "reopen", "by": args.by, "reason": args.reason}, ensure_ascii=False) + "\n")
    upsert_ledger(root.parent, {"案件フォルダ": root.name, "状態": "計画差戻し", "最終更新": args.today})
    dump({"plan_status": "draft", "by": args.by, "reason": args.reason})


def cmd_status(args) -> None:
    root = engagement_dir(args.engagement)
    # 案件台帳は親フォルダにあるので、案件フォルダの中にも残す。state.py はこちらを読む
    with open(root / "06_ログ" / "承認記録.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps({"ts": args.today, "action": "status", "status": args.status}, ensure_ascii=False) + "\n")
    name = upsert_ledger(root.parent, {"案件フォルダ": root.name, "状態": args.status, "最終更新": args.today})
    dump({"ledger": name, "status": args.status})


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("command", choices=["new", "approve", "reopen", "status"])
    ap.add_argument("--root", help="new: 案件フォルダを並べる親フォルダ")
    ap.add_argument("--client")
    ap.add_argument("--period-end", help="決算期末日 YYYY-MM-DD")
    ap.add_argument("--engagement-type")
    ap.add_argument("--prior", help="前期の案件フォルダ")
    ap.add_argument("--engagement", help="対象の案件フォルダ")
    ap.add_argument("--by")
    ap.add_argument("--reason")
    ap.add_argument("--status")
    ap.add_argument("--today", required=True)
    args = ap.parse_args()
    if args.command == "new":
        if not (args.root and args.client and args.period_end):
            sys.exit("new には --root --client --period-end が必要です")
        cmd_new(args)
    else:
        if not args.engagement:
            sys.exit("--engagement が必要です")
        if args.command in ("approve", "reopen") and not args.by:
            sys.exit("--by（承認者・差戻し者の名前）が必要です")
        if args.command == "status" and not args.status:
            sys.exit("--status が必要です")
        {"approve": cmd_approve, "reopen": cmd_reopen, "status": cmd_status}[args.command](args)
    return 0


if __name__ == "__main__":
    sys.exit(main())
