#!/usr/bin/env python3
"""資料依頼（PBC）キューの管理。依頼・送信・督促・受領・充足判定を1つの記録で追う。

記録は 04_依頼管理/pbc.jsonl に追記のみで残し（イベントログ）、現在の状態は再生して求める。
人が見る表は `report` で 04_依頼管理/依頼トラッカー.xlsx に書き出す。

サブコマンド:
  add     依頼を追加。同じ資料（資料種別・期間・勘定が同じ）の未完了依頼があれば統合する。
          書き方が違う同じ資料は "same_as": "PBC0003" で既存の依頼に統合する
  list    現在の状態を一覧（--status で絞込）
  mark    状態を更新（drafted/sent/reminded/received/accepted/rejected/escalated/cancelled）
  due     --today 時点で、送るべき依頼・督促すべき依頼・エスカレーションすべき依頼を返す
  report  依頼トラッカー.xlsx を作り直す

例:
  python3 pbc.py add --engagement . --today 2026-10-27 --json '{"area":"売上","doc_type":"得意先別売掛金残高表",
      "period":"2026-03-31時点","accounts":["1130"],"procedure":"残高確認の母集団","reason":"MUSの母集団に使う","to":"keiri"}'
  python3 pbc.py add --engagement . --today 2026-10-28 --json '{"same_as":"PBC0001","area":"売掛金","doc_type":"得意先別の売掛金残高",
      "period":"2026-03-31","accounts":["1130"],"procedure":"AR-01","reason":"残高の照合","to":"keiri"}'
  python3 pbc.py mark --engagement . --id PBC0003 --status sent --today 2026-10-27 --note "メールID xxx"
  python3 pbc.py due --engagement . --today 2026-10-30
"""
import argparse
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import dump, load_engagement  # noqa: E402

STATUSES = ["queued", "drafted", "sent", "reminded", "received", "accepted", "rejected", "escalated", "cancelled"]
OPEN = {"queued", "drafted", "sent", "reminded", "received", "rejected"}
REQUIRED = ["area", "doc_type", "period", "reason", "to"]


def norm(s: str) -> str:
    return re.sub(r"[\s　・、,（）()]+", "", str(s)).lower()


def key_of(req: dict) -> str:
    return "|".join([norm(req["doc_type"]), norm(req["period"]), ",".join(sorted(map(str, req.get("accounts", []))))])


def log_path(root: Path) -> Path:
    return root / "04_依頼管理" / "pbc.jsonl"


def append(root: Path, event: dict) -> None:
    p = log_path(root)
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "a", encoding="utf-8") as f:
        f.write(json.dumps(event, ensure_ascii=False) + "\n")


def replay(root: Path) -> dict[str, dict]:
    state: dict[str, dict] = {}
    p = log_path(root)
    if not p.exists():
        return state
    for line in p.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        e = json.loads(line)
        if e["event"] == "created":
            r = dict(e["request"])
            r.update(id=e["id"], status="queued", created=e["ts"], areas=[r["area"]],
                     procedures=[r.get("procedure", "")], reasons=[r["reason"]],
                     history=[], reminders=0, last_contact=None, evidence=[])
            state[e["id"]] = r
        elif e["event"] == "merged":
            r = state[e["id"]]
            if e["request"]["area"] not in r["areas"]:
                r["areas"].append(e["request"]["area"])
            r["procedures"].append(e["request"].get("procedure", ""))
            r["reasons"].append(e["request"]["reason"])
        elif e["event"] == "status":
            r = state[e["id"]]
            r["status"] = e["status"]
            r["history"].append({"ts": e["ts"], "status": e["status"], "note": e.get("note", "")})
            if e["status"] in ("sent", "reminded"):
                r["last_contact"] = e["ts"]
            if e["status"] == "reminded":
                r["reminders"] += 1
            if e.get("evidence"):
                r["evidence"].extend(e["evidence"])
    return state


def business_days_between(a: date, b: date, holidays: set[date]) -> int:
    n, d = 0, a
    while d < b:
        d += timedelta(days=1)
        if d.weekday() < 5 and d not in holidays:
            n += 1
    return n


def cmd_add(root: Path, args) -> None:
    req = json.loads(args.json)
    missing = [k for k in REQUIRED if not req.get(k)]
    if missing:
        sys.exit(f"依頼に必要な項目がありません: {', '.join(missing)}")
    eng = load_engagement(root)
    contacts = {c["key"] for c in eng.get("contacts", [])}
    if req["to"] not in contacts:
        sys.exit(f"宛先キー {req['to']!r} は engagement.json の contacts にありません。計画承認で登録された担当者だけに依頼できます")
    state = replay(root)
    same_as = req.pop("same_as", None)
    if same_as:
        target = state.get(same_as)
        if target is None:
            sys.exit(f"same_as の {same_as} がありません")
        if target["status"] not in OPEN:
            sys.exit(f"{same_as} は {target['status']} で完了しています。新しく依頼するなら same_as を外してください")
        append(root, {"ts": args.today, "event": "merged", "id": same_as, "request": req})
        dump({"action": "merged", "id": same_as, "areas": sorted(set(target["areas"] + [req["area"]]))})
        return
    k = key_of(req)
    for r in state.values():
        if r["status"] in OPEN and key_of(r) == k:
            append(root, {"ts": args.today, "event": "merged", "id": r["id"], "request": req})
            dump({"action": "merged", "id": r["id"], "areas": sorted(set(r["areas"] + [req["area"]]))})
            return
    new_id = f"PBC{len(state) + 1:04d}"
    append(root, {"ts": args.today, "event": "created", "id": new_id, "request": req})
    dump({"action": "created", "id": new_id})


def cmd_mark(root: Path, args) -> None:
    if args.status not in STATUSES:
        sys.exit(f"status は {', '.join(STATUSES)} のいずれか")
    state = replay(root)
    if args.id not in state:
        sys.exit(f"{args.id} がありません")
    ev = {"ts": args.today, "event": "status", "id": args.id, "status": args.status, "note": args.note or ""}
    if args.evidence:
        ev["evidence"] = args.evidence
    append(root, ev)
    dump({"id": args.id, "status": args.status})


def cmd_list(root: Path, args) -> None:
    state = replay(root)
    rows = [r for r in state.values() if not args.status or r["status"] in args.status]
    dump(rows)


def due(root: Path, today_s: str) -> dict:
    """today_s 時点で、送るべき依頼・督促すべき依頼・エスカレーションすべき依頼を返す。state.py からも使う。"""
    eng = load_engagement(root)
    rules = eng.get("pbc_rules", {})
    remind_after = int(rules.get("remind_after_business_days", 3))
    escalate_after = int(rules.get("escalate_after_business_days", 3))
    max_reminders = int(rules.get("max_reminders", 1))
    holidays = {date.fromisoformat(h) for h in eng.get("holidays", [])}
    deadline = eng.get("evidence_deadline")
    today = date.fromisoformat(today_s)

    to_send: dict[str, list] = {}
    remind, escalate, past_deadline = [], [], []
    for r in replay(root).values():
        if r["status"] in ("queued", "drafted"):
            to_send.setdefault(r["to"], []).append(r["id"])
        elif r["status"] in ("sent", "reminded"):
            since = business_days_between(date.fromisoformat(r["last_contact"]), today, holidays)
            if r["reminders"] < max_reminders and since >= remind_after:
                remind.append({"id": r["id"], "to": r["to"], "business_days_waiting": since})
            elif r["reminders"] >= max_reminders and since >= escalate_after:
                escalate.append({"id": r["id"], "to": r["to"], "business_days_waiting": since})
        if deadline and r["status"] in OPEN and today > date.fromisoformat(deadline):
            past_deadline.append(r["id"])
    return {
        "today": today_s,
        "send_batches": [{"to": k, "ids": v} for k, v in to_send.items()],
        "remind": remind,
        "escalate": escalate,
        "past_evidence_deadline": past_deadline,
    }


def cmd_due(root: Path, args) -> None:
    dump(due(root, args.today))


def cmd_report(root: Path, args) -> None:
    state = replay(root)
    eng = load_engagement(root)
    names = {c["key"]: f'{c.get("role", "")} {c.get("name", "")}'.strip() for c in eng.get("contacts", [])}
    header = ["依頼ID", "状態", "領域", "資料種別", "期間", "対象勘定", "使う手続", "理由", "依頼先",
              "登録日", "最終連絡日", "督促回数", "受領証拠ID", "履歴"]
    rows = []
    for r in state.values():
        rows.append([
            r["id"], r["status"], "、".join(r["areas"]), r["doc_type"], r["period"],
            "、".join(map(str, r.get("accounts", []))), " / ".join(p for p in r["procedures"] if p),
            " / ".join(r["reasons"]), names.get(r["to"], r["to"]), r["created"], r["last_contact"] or "",
            r["reminders"], "、".join(r["evidence"]),
            " → ".join(f'{h["ts"]} {h["status"]}' for h in r["history"]),
        ])
    out_dir = root / "04_依頼管理"
    try:
        import openpyxl  # noqa: PLC0415
        from openpyxl.styles import Alignment, Font  # noqa: PLC0415
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "依頼トラッカー"
        ws.append(header)
        for row in rows:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
        widths = [10, 10, 16, 24, 16, 12, 24, 36, 18, 12, 12, 8, 14, 40]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        path = out_dir / "依頼トラッカー.xlsx"
        wb.save(path)
    except ImportError:
        import csv  # noqa: PLC0415
        path = out_dir / "依頼トラッカー.csv"
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)
    dump({"written": str(path.relative_to(root)), "requests": len(rows)})


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("command", choices=["add", "mark", "list", "due", "report"])
    ap.add_argument("--engagement", required=True)
    ap.add_argument("--today", help="YYYY-MM-DD。add/mark/due で必須")
    ap.add_argument("--json")
    ap.add_argument("--id")
    ap.add_argument("--status", nargs="*")
    ap.add_argument("--note")
    ap.add_argument("--evidence", nargs="*")
    args = ap.parse_args()
    root = Path(args.engagement)
    if args.command in ("add", "mark", "due") and not args.today:
        sys.exit("--today を指定してください")
    if args.command == "mark":
        if not args.status or len(args.status) != 1 or not args.id:
            sys.exit("mark には --id と --status（1つ）が必要です")
        args.status = args.status[0]
    {"add": cmd_add, "mark": cmd_mark, "list": cmd_list, "due": cmd_due, "report": cmd_report}[args.command](root, args)
    return 0


if __name__ == "__main__":
    sys.exit(main())
